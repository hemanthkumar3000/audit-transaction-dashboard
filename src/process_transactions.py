"""
Transaction Data Cleaning, Audit Exception Flagging & SQLite Pipeline
Author: Audit Analytics Team
Description: Cleans transaction data, applies audit exception logic, loads to SQLite,
and exports a formatted Excel exception report.
"""

import os
import sqlite3
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def process_audit_data():
    raw_csv_path = os.path.join("data", "transactions.csv")
    output_excel_path = os.path.join("output", "audit_exception_report.xlsx")
    db_path = os.path.join("data", "audit_transactions.db")

    print(f"Loading raw transaction data from {raw_csv_path}...")
    df = pd.read_csv(raw_csv_path)

    # 1. Data Cleaning
    print("Performing data type conversions and validation...")
    df["Date"] = pd.to_datetime(df["Date"])
    df["Amount"] = pd.to_numeric(df["Amount"], errors="coerce")

    # Drop any nulls if present
    df = df.dropna(subset=["Transaction_ID", "Date", "Amount"]).copy()

    # Sort deterministically
    df = df.sort_values(by=["Date", "Transaction_ID"]).reset_index(drop=True)

    # 2. Audit Exception Rules
    print("Evaluating audit exception rules...")
    # Rule 1: High-value transaction (> ₹50,000)
    df["high_value_flag"] = df["Amount"] > 50000

    # Rule 2: Weekend transaction (Saturday or Sunday)
    # pandas dt.dayofweek: Monday=0 ... Saturday=5, Sunday=6
    df["weekend_flag"] = df["Date"].dt.dayofweek >= 5

    # Rule 3: Possible Duplicate (same Date, Vendor, Amount)
    df["duplicate_flag"] = df.duplicated(
        subset=["Date", "Vendor", "Amount"],
        keep=False
    )

    # Combine flags into a consolidated reason string
    df["flag_reason"] = ""
    df.loc[df["high_value_flag"], "flag_reason"] += "High Value; "
    df.loc[df["weekend_flag"], "flag_reason"] += "Weekend; "
    df.loc[df["duplicate_flag"], "flag_reason"] += "Possible Duplicate; "
    df["flag_reason"] = df["flag_reason"].str.rstrip("; ")

    # Flag status indicator
    df["is_flagged"] = df["flag_reason"] != ""

    # Day of week name for reporting clarity
    df["Day_Of_Week"] = df["Date"].dt.day_name()
    # Format Date as string YYYY-MM-DD for clean DB & Excel display
    df["Date_Str"] = df["Date"].dt.strftime("%Y-%m-%d")

    # 3. Load into SQLite Database
    print(f"Loading data into SQLite database at {db_path}...")
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()

    # Prepare table schema matching raw transactions
    cursor.execute("DROP TABLE IF EXISTS transactions;")
    cursor.execute("""
    CREATE TABLE transactions (
        Transaction_ID TEXT PRIMARY KEY,
        Date TEXT NOT NULL,
        Department TEXT NOT NULL,
        Employee_ID TEXT NOT NULL,
        Vendor TEXT NOT NULL,
        Amount REAL NOT NULL,
        Payment_Method TEXT NOT NULL
    );
    """)

    # Insert cleaned records into transactions table
    db_records = df[[
        "Transaction_ID", "Date_Str", "Department",
        "Employee_ID", "Vendor", "Amount", "Payment_Method"
    ]].rename(columns={"Date_Str": "Date"})

    db_records.to_sql("transactions", conn, if_exists="append", index=False)

    # Also store the audit flagged table in SQLite for audit querying
    cursor.execute("DROP TABLE IF EXISTS audit_exceptions;")
    audit_db_records = df[df["is_flagged"]][[
        "Transaction_ID", "Date_Str", "Department",
        "Employee_ID", "Vendor", "Amount", "Payment_Method",
        "high_value_flag", "weekend_flag", "duplicate_flag", "flag_reason"
    ]].rename(columns={"Date_Str": "Date"})

    audit_db_records.to_sql("audit_exceptions", conn, if_exists="append", index=False)
    conn.commit()
    conn.close()
    print("Data successfully committed to SQLite.")

    # 4. Generate Audit Summary Statistics
    total_txns = len(df)
    total_amt = df["Amount"].sum()
    flagged_df = df[df["is_flagged"]].copy()
    flagged_txns = len(flagged_df)
    flagged_amt = flagged_df["Amount"].sum()
    pct_flagged_count = (flagged_txns / total_txns) * 100
    pct_flagged_amt = (flagged_amt / total_amt) * 100

    high_val_count = df["high_value_flag"].sum()
    weekend_count = df["weekend_flag"].sum()
    duplicate_count = df["duplicate_flag"].sum()

    print("\n--- AUDIT SUMMARY METRICS ---")
    print(f"Total Transactions:        {total_txns:,}")
    print(f"Total Amount:              INR {total_amt:,.2f}")
    print(f"Flagged Transactions:      {flagged_txns:,} ({pct_flagged_count:.1f}%)")
    print(f"Flagged Amount:            INR {flagged_amt:,.2f} ({pct_flagged_amt:.1f}%)")
    print(f" - High-Value (> 50k):     {high_val_count:,}")
    print(f" - Weekend Postings:       {weekend_count:,}")
    print(f" - Possible Duplicates:    {duplicate_count:,}")

    # 5. Export Multi-Sheet Professional Excel Report
    print(f"\nExporting audit exception report to {output_excel_path}...")
    wb = openpyxl.Workbook()

    # Style definitions
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    sub_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    alert_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    kpi_card_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    font_title = Font(name="Segoe UI", size=16, bold=True, color="1F4E79")
    font_section = Font(name="Segoe UI", size=12, bold=True, color="1F4E79")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Segoe UI", size=10, bold=True)
    font_regular = Font(name="Segoe UI", size=10)
    font_kpi_value = Font(name="Segoe UI", size=14, bold=True, color="C00000")

    border_thin = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3")
    )
    border_total = Border(
        top=Side(style="thin", color="000000"),
        bottom=Side(style="double", color="000000")
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # Sheet 1: Executive Summary
    ws_sum = wb.active
    ws_sum.title = "Executive Summary"
    ws_sum.views.sheetView[0].showGridLines = True

    ws_sum["A1"] = "AUDIT TRANSACTION EXCEPTION REPORT"
    ws_sum["A1"].font = font_title
    ws_sum["A2"] = "Automated Review & Risk Flagging Analysis"
    ws_sum["A2"].font = Font(name="Segoe UI", size=11, italic=True, color="595959")

    # KPI Summary Cards
    ws_sum["A4"] = "Audit Overview KPI Metrics"
    ws_sum["A4"].font = font_section

    kpi_headers = ["Metric Description", "Value", "% of Portfolio", "Audit Impact / Action"]
    for col_idx, h in enumerate(kpi_headers, 1):
        cell = ws_sum.cell(row=5, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = align_center

    kpi_rows = [
        ("Total Transactions Analyzed", total_txns, "100.0%", "Complete dataset scoped"),
        ("Total Transaction Volume (₹)", total_amt, "100.0%", "Total expenditure audited"),
        ("Total Flagged Exceptions", flagged_txns, f"{pct_flagged_count:.1f}%", "Requires senior auditor review"),
        ("Total Flagged Amount (₹)", flagged_amt, f"{pct_flagged_amt:.1f}%", "Substantial financial exposure"),
        ("High-Value Transactions (> ₹50,000)", high_val_count, f"{(high_val_count/total_txns)*100:.1f}%", "Check management authorization"),
        ("Weekend Transactions (Sat/Sun)", weekend_count, f"{(weekend_count/total_txns)*100:.1f}%", "Validate after-hours logging"),
        ("Duplicate-Looking Records", duplicate_count, f"{(duplicate_count/total_txns)*100:.1f}%", "Potential duplicate billing / AP error")
    ]

    for row_idx, (desc, val, pct, impact) in enumerate(kpi_rows, 6):
        c1 = ws_sum.cell(row=row_idx, column=1, value=desc)
        c2 = ws_sum.cell(row=row_idx, column=2, value=val)
        c3 = ws_sum.cell(row=row_idx, column=3, value=pct)
        c4 = ws_sum.cell(row=row_idx, column=4, value=impact)

        c1.font = font_bold if "Total" in desc else font_regular
        c2.font = font_kpi_value if "Flagged" in desc else font_bold
        c3.font = font_regular
        c4.font = font_regular

        c1.alignment = align_left
        c2.alignment = align_right
        c3.alignment = align_center
        c4.alignment = align_left

        for c in (c1, c2, c3, c4):
            c.border = border_thin

        if "₹" in desc or isinstance(val, (int, float)) and val > 1000:
            c2.number_format = "₹#,##0.00"
        elif isinstance(val, int):
            c2.number_format = "#,##0"

    # Department Breakdown Table on Summary Sheet
    dept_start_row = 15
    ws_sum.cell(row=dept_start_row, column=1, value="Department Risk Breakdown").font = font_section

    dept_cols = ["Department", "Total Txns", "Flagged Txns", "Flagged %", "Total Spend (₹)", "Flagged Spend (₹)"]
    for col_idx, col_name in enumerate(dept_cols, 1):
        cell = ws_sum.cell(row=dept_start_row + 1, column=col_idx, value=col_name)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = align_center

    dept_summary = df.groupby("Department").agg(
        Total_Txns=("Transaction_ID", "count"),
        Flagged_Txns=("is_flagged", "sum"),
        Total_Amount=("Amount", "sum"),
        Flagged_Amount=("Amount", lambda x: x[df.loc[x.index, "is_flagged"]].sum())
    ).reset_index()

    dept_summary["Flagged_Pct"] = (dept_summary["Flagged_Txns"] / dept_summary["Total_Txns"]) * 100
    dept_summary = dept_summary.sort_values(by="Flagged_Amount", ascending=False)

    curr_row = dept_start_row + 2
    for _, row in dept_summary.iterrows():
        c1 = ws_sum.cell(row=curr_row, column=1, value=row["Department"])
        c2 = ws_sum.cell(row=curr_row, column=2, value=int(row["Total_Txns"]))
        c3 = ws_sum.cell(row=curr_row, column=3, value=int(row["Flagged_Txns"]))
        c4 = ws_sum.cell(row=curr_row, column=4, value=row["Flagged_Pct"] / 100)
        c5 = ws_sum.cell(row=curr_row, column=5, value=row["Total_Amount"])
        c6 = ws_sum.cell(row=curr_row, column=6, value=row["Flagged_Amount"])

        c1.alignment = align_left
        c2.alignment = align_right
        c3.alignment = align_right
        c4.alignment = align_right
        c5.alignment = align_right
        c6.alignment = align_right

        c2.number_format = "#,##0"
        c3.number_format = "#,##0"
        c4.number_format = "0.0%"
        c5.number_format = "₹#,##0.00"
        c6.number_format = "₹#,##0.00"

        for c in (c1, c2, c3, c4, c5, c6):
            c.font = font_regular
            c.border = border_thin
        curr_row += 1

    # Total row for department summary
    tot_row = curr_row
    ws_sum.cell(row=tot_row, column=1, value="Total Portfolio").font = font_bold
    ws_sum.cell(row=tot_row, column=2, value=f"=SUM(B{dept_start_row+2}:B{tot_row-1})").font = font_bold
    ws_sum.cell(row=tot_row, column=3, value=f"=SUM(C{dept_start_row+2}:C{tot_row-1})").font = font_bold
    ws_sum.cell(row=tot_row, column=4, value=f"=C{tot_row}/B{tot_row}").font = font_bold
    ws_sum.cell(row=tot_row, column=5, value=f"=SUM(E{dept_start_row+2}:E{tot_row-1})").font = font_bold
    ws_sum.cell(row=tot_row, column=6, value=f"=SUM(F{dept_start_row+2}:F{tot_row-1})").font = font_bold

    ws_sum.cell(row=tot_row, column=2).number_format = "#,##0"
    ws_sum.cell(row=tot_row, column=3).number_format = "#,##0"
    ws_sum.cell(row=tot_row, column=4).number_format = "0.0%"
    ws_sum.cell(row=tot_row, column=5).number_format = "₹#,##0.00"
    ws_sum.cell(row=tot_row, column=6).number_format = "₹#,##0.00"

    for c in range(1, 7):
        cell = ws_sum.cell(row=tot_row, column=c)
        cell.border = border_total
        cell.fill = sub_fill

    # Sheet 2: Flagged Exceptions Detailed Table
    ws_exc = wb.create_sheet(title="Audit Exceptions")
    ws_exc.views.sheetView[0].showGridLines = True

    exc_headers = [
        "Transaction_ID", "Date", "Day_Of_Week", "Department", "Employee_ID",
        "Vendor", "Amount (₹)", "Payment_Method",
        "High Value", "Weekend", "Duplicate", "Flag Reason"
    ]

    for col_idx, h in enumerate(exc_headers, 1):
        cell = ws_exc.cell(row=1, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = align_center

    # Sort flagged records by Amount descending
    flagged_sorted = flagged_df.sort_values(by="Amount", ascending=False)
    for r_idx, (_, r) in enumerate(flagged_sorted.iterrows(), 2):
        ws_exc.cell(row=r_idx, column=1, value=r["Transaction_ID"]).alignment = align_center
        ws_exc.cell(row=r_idx, column=2, value=r["Date_Str"]).alignment = align_center
        ws_exc.cell(row=r_idx, column=3, value=r["Day_Of_Week"]).alignment = align_center
        ws_exc.cell(row=r_idx, column=4, value=r["Department"]).alignment = align_left
        ws_exc.cell(row=r_idx, column=5, value=r["Employee_ID"]).alignment = align_center
        ws_exc.cell(row=r_idx, column=6, value=r["Vendor"]).alignment = align_left

        amt_cell = ws_exc.cell(row=r_idx, column=7, value=r["Amount"])
        amt_cell.alignment = align_right
        amt_cell.number_format = "₹#,##0.00"

        ws_exc.cell(row=r_idx, column=8, value=r["Payment_Method"]).alignment = align_center
        ws_exc.cell(row=r_idx, column=9, value="YES" if r["high_value_flag"] else "NO").alignment = align_center
        ws_exc.cell(row=r_idx, column=10, value="YES" if r["weekend_flag"] else "NO").alignment = align_center
        ws_exc.cell(row=r_idx, column=11, value="YES" if r["duplicate_flag"] else "NO").alignment = align_center

        reason_cell = ws_exc.cell(row=r_idx, column=12, value=r["flag_reason"])
        reason_cell.alignment = align_left

        # Soft highlight for rows with multiple exception rules
        flag_count = int(r["high_value_flag"]) + int(r["weekend_flag"]) + int(r["duplicate_flag"])
        row_fill = alert_fill if flag_count > 1 else None

        for col_idx in range(1, 13):
            cell = ws_exc.cell(row=r_idx, column=col_idx)
            cell.font = font_regular
            cell.border = border_thin
            if row_fill:
                cell.fill = row_fill

    # Sheet 3: All Cleaned Transactions
    ws_all = wb.create_sheet(title="All Cleaned Transactions")
    ws_all.views.sheetView[0].showGridLines = True

    all_headers = [
        "Transaction_ID", "Date", "Department", "Employee_ID",
        "Vendor", "Amount (₹)", "Payment_Method", "Is_Flagged", "Flag_Reason"
    ]
    for col_idx, h in enumerate(all_headers, 1):
        cell = ws_all.cell(row=1, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = PatternFill(start_color="333F48", end_color="333F48", fill_type="solid")
        cell.alignment = align_center

    for r_idx, (_, r) in enumerate(df.iterrows(), 2):
        ws_all.cell(row=r_idx, column=1, value=r["Transaction_ID"]).alignment = align_center
        ws_all.cell(row=r_idx, column=2, value=r["Date_Str"]).alignment = align_center
        ws_all.cell(row=r_idx, column=3, value=r["Department"]).alignment = align_left
        ws_all.cell(row=r_idx, column=4, value=r["Employee_ID"]).alignment = align_center
        ws_all.cell(row=r_idx, column=5, value=r["Vendor"]).alignment = align_left

        amt_cell = ws_all.cell(row=r_idx, column=6, value=r["Amount"])
        amt_cell.alignment = align_right
        amt_cell.number_format = "₹#,##0.00"

        ws_all.cell(row=r_idx, column=7, value=r["Payment_Method"]).alignment = align_center
        ws_all.cell(row=r_idx, column=8, value="FLAGGED" if r["is_flagged"] else "OK").alignment = align_center
        ws_all.cell(row=r_idx, column=9, value=r["flag_reason"] if r["is_flagged"] else "-").alignment = align_left

        for col_idx in range(1, 10):
            cell = ws_all.cell(row=r_idx, column=col_idx)
            cell.font = font_regular
            cell.border = border_thin

    # Auto-fit column widths across all sheets
    for ws in [ws_sum, ws_exc, ws_all]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = cell.value
                if val:
                    val_str = str(val)
                    if len(val_str) > max_len:
                        max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(output_excel_path)
    print(f"Excel report successfully written to {output_excel_path}")

    return df

if __name__ == "__main__":
    process_audit_data()
